from __future__ import annotations

import pandas as pd
import pytest
from openpyxl import Workbook

from pbi_xbrl.excel_writer_placeholders import (
    DEFAULT_EMPTY_SHEET_MESSAGE,
    write_empty_sheet_placeholder,
)
from pbi_xbrl.market_data.contracts import (
    MARKET_ROWS_DF_NORMALIZED_ATTR,
    normalize_market_rows_df,
    require_market_columns,
)
from pbi_xbrl.qa_outputs import dedupe_audit_like_rows, rows_to_audit_like_frame


def test_write_empty_sheet_placeholder_uses_current_build_message() -> None:
    wb = Workbook()
    ws = wb.active
    write_empty_sheet_placeholder(ws)
    assert str(ws["A1"].value or "").strip() == DEFAULT_EMPTY_SHEET_MESSAGE


def test_market_rows_contract_normalizes_missing_columns_and_marks_frame() -> None:
    df = normalize_market_rows_df([{"series_key": "ethanol_demo", "price_value": "1.23"}])
    assert bool(df.attrs.get(MARKET_ROWS_DF_NORMALIZED_ATTR))
    assert "observation_date" in df.columns
    assert "quarter" in df.columns
    assert "contract_tenor" in df.columns
    assert pd.isna(df.loc[0, "observation_date"])
    assert float(df.loc[0, "price_value"]) == pytest.approx(1.23)


def test_market_rows_contract_require_columns_fails_fast_for_internal_paths() -> None:
    with pytest.raises(ValueError, match="requires normalized market columns: observation_date"):
        require_market_columns(
            pd.DataFrame({"series_key": ["ethanol_demo"]}),
            ["series_key", "observation_date"],
            contract_name="demo_contract",
        )


def test_dedupe_audit_like_rows_preserves_first_seen_order() -> None:
    frame = rows_to_audit_like_frame(
        [
            {"quarter": "2025-12-31", "metric": "QA_QTR", "severity": "warn", "message": "same", "source": "pipeline"},
            {"quarter": "2025-12-31", "metric": "QA_QTR", "severity": "warn", "message": "same", "source": "pipeline"},
            {"quarter": "2025-12-31", "metric": "QA_QTR", "severity": "warn", "message": "other", "source": "pipeline"},
        ],
        defaults={
            "quarter": pd.NaT,
            "metric": "",
            "severity": "info",
            "message": "",
            "source": "",
        },
    )
    deduped = dedupe_audit_like_rows(frame)
    assert len(deduped) == 2
    assert str(deduped.iloc[0]["message"]) == "same"
    assert str(deduped.iloc[1]["message"]) == "other"
